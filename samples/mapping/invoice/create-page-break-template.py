"""Generate the nested, explicitly paginated invoice sample from template.xlsx."""
from copy import copy
from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

root = Path(__file__).parent
source = load_workbook(root / 'template.xlsx').active
sheet = Workbook().active
sheet.title = source.title
for key, dimension in source.column_dimensions.items():
    sheet.column_dimensions[key] = copy(dimension)
for row in source.iter_rows(min_row=1, max_row=21, max_col=6):
    for original in row:
        cell = sheet.cell(original.row + 1, original.column, original.value)
        cell.font = copy(original.font)
        cell.fill = copy(original.fill)
        cell.border = copy(original.border)
        cell.alignment = copy(original.alignment)
        cell.number_format = original.number_format
    sheet.row_dimensions[row[0].row + 1].height = source.row_dimensions[row[0].row].height
for area in source.merged_cells.ranges:
    if str(area) != 'A1:F1':
        sheet.merge_cells(start_row=area.min_row + 1, end_row=area.max_row + 1,
                          start_column=area.min_col, end_column=area.max_col)
sheet.merge_cells('B2:F2')
sheet['A1'] = '**@start-array $.pages[*] as page'
sheet['A2'] = '**@page-break'
sheet['B2'] = '**@page.title'
sheet['B2'].font = copy(source['A1'].font)
sheet['B2'].fill = copy(source['A1'].fill)
sheet['B2'].alignment = Alignment(horizontal='center', vertical='center')
sheet['A10'] = '**@start-array @page.items[*] as item'
sheet['F11'] = '**@item.lineTotal'
sheet['A23'] = '**@end-array'
for address, expression in [('F7', '**@page.totals.subtotal'), ('F14', '**@page.totals.subtotal'),
                            ('F15', '**@page.totals.tax'), ('F16', '**@page.totals.total'),
                            ('B19', '**@page.totals.total | format("N2", "de-DE")'),
                            ('B22', '**@page.totals.total | format("C0", "ja-JP")')]:
    sheet[address] = expression
sheet.sheet_view.showGridLines = False
# Anchor the print area beyond the repeated block so all copies are included.
sheet.row_dimensions[24].height = 12
sheet.print_area = 'A1:F24'
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_setup = copy(source.page_setup)
sheet.page_margins = copy(source.page_margins)
sheet.parent.save(root / 'page-breaks' / 'template.xlsx')
data = json.loads((root / 'data.json').read_text())
items = data.pop('items')
data.pop('totals')
data['pages'] = []
for i, group in enumerate([items[:2], [], items[2:]]):
    subtotal = sum(item['quantity'] * item['price'] for item in group)
    tax = sum(item['quantity'] * item['price'] * item['taxRate'] for item in group)
    data['pages'].append({'title': f'請求書 明細 {i + 1}/3', 'items': group,
                          'totals': {'subtotal': round(subtotal, 2), 'tax': round(tax, 2), 'total': round(subtotal + tax, 2)}})
(root / 'page-breaks' / 'data.json').write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')
