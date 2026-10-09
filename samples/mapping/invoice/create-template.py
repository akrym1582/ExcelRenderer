"""Regenerate template.xlsx with Python 3 and openpyxl (pip install openpyxl)."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

sheet = Workbook().active
sheet.title = "請求書"
for column, width in zip("ABCDEF", (12, 19, 13, 10, 14, 19)):
    sheet.column_dimensions[column].width = width
for row in range(1, 23):
    sheet.row_dimensions[row].height = 25
    for cell in sheet[row][:6]:
        cell.font = Font(name="Noto Sans JP", size=10, color="FF243447")
        cell.alignment = Alignment(vertical="center")
sheet.row_dimensions[1].height = 40
sheet.row_dimensions[7].height = 12
sheet.row_dimensions[16].height = 12
values = {
    "A1": "請 求 書", "A2": "請求先", "B2": "**$.customer.name", "E2": "請求番号", "F2": "**invoice.number",
    "A3": "住所", "B3": "**$['customer']['address.line']", "E3": "発行日", "F3": '**invoice.issuedAt | date("yyyy/MM/dd")',
    "A4": "担当者", "B4": "**.customer.contacts[0].name", "E4": "通貨", "F4": '**$["meta|info"].currency',
    "A5": "元の日付", "B5": "**invoice.issuedAt", "E5": "入金済", "F5": "**invoice.paid",
    "A6": "作業時間", "B6": "**invoice.duration", "E6": "税抜合計", "F6": "**totals.subtotal",
    "A8": "コード", "B8": "品名 / 備考", "D8": "数量", "E8": "単価", "F8": "金額",
    "A9": "**@start-array $.items[*] as item",
    "A10": '**@item.code | format("0000")', "B10": "**@item.name", "D10": "**@item.quantity", "E10": "**@item.price",
    "A11": "備考", "B11": "**@item.note", "E11": "税率", "F11": "**@item.taxRate",
    "A12": "**@end-array",
    "E13": "小計", "F13": "**totals.subtotal", "E14": "消費税", "F14": "**totals.tax", "E15": "ご請求額", "F15": "**totals.total",
    "A17": "表示形式の確認",
    "A18": "N2 (de-DE)", "B18": '**totals.total | format("N2", "de-DE")', "E18": "支払期限", "F18": '**invoice.dueAt | date("dd MMM yyyy", "en-US")',
    "A19": "ゼロ", "B19": "**checks.zero", "E19": "指数", "F19": "**checks.scientific",
    "A20": r"\**literal", "B20": "**checks.negative", "E20": "空欄 (null)", "F20": "**checks.empty",
    "A21": "通貨文字列", "B21": '**totals.total | format("C0", "ja-JP")', "E21": "文字列番号", "F21": "**checks.textCode",
}
for address, value in values.items():
    sheet[address] = value
sheet["F10"] = "=D10*E10"
for area in ("A1:F1", "B2:D2", "B3:D3", "B4:D4", "B5:D5", "B6:D6", "B8:C8", "B10:C10", "B11:D11", "A17:F17", "B18:C18", "B21:D21"):
    sheet.merge_cells(area)
for row in (1, 8, 17):
    for cell in sheet[row][:6]:
        cell.fill = PatternFill("solid", fgColor="FF243447")
        cell.font = Font(name="Noto Sans JP", size=20 if row == 1 else 10, bold=True, color="FFFFFFFF")
        cell.alignment = Alignment(vertical="center", horizontal="center" if row == 1 else "left")
line = Side(style="thin", color="FFCBD5E1")
for row in (10, 11):
    for cell in sheet[row][:6]:
        cell.border = Border(bottom=line)
        if row == 11:
            cell.fill = PatternFill("solid", fgColor="FFF1F5F9")
sheet.row_dimensions[10].height = 28
sheet.row_dimensions[11].height = 24
for address in ("F6", "E10", "F10", "F13", "F14", "F15"):
    sheet[address].number_format = '"¥"#,##0.00;[Red]-"¥"#,##0.00;"¥"0.00'
for address in ("D10", "B19", "B20"):
    sheet[address].number_format = '#,##0.00;[Red](#,##0.00);"-"'
sheet["F11"].number_format = "0.0%"
sheet["F19"].number_format = "0.00E+00"
sheet["B5"].number_format = "yyyy/mm/dd"
sheet["B6"].number_format = "[h]:mm:ss"
for row in (13, 14, 15):
    for cell in sheet[row][4:6]:
        cell.fill = PatternFill("solid", fgColor="FFE2E8F0")
        cell.font = Font(name="Noto Sans JP", size=12 if row == 15 else 10, bold=True, color="FF243447")
for row in range(2, 22):
    for column in (4, 5, 6):
        sheet.cell(row, column).alignment = Alignment(vertical="center", horizontal="right")
sheet.sheet_view.showGridLines = False
sheet.print_options.gridLines = False
sheet.print_area = "A1:F21"
sheet.sheet_properties.pageSetUpPr.fitToPage = True
sheet.page_setup.orientation = "portrait"
sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
sheet.page_setup.fitToWidth = 1
sheet.page_setup.fitToHeight = 1
sheet.page_margins = PageMargins(left=0.25, right=0.25, top=0.3, bottom=0.3, header=0, footer=0)
sheet.parent.save(Path(__file__).with_name("template.xlsx"))
