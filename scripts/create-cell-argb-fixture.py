#!/usr/bin/env python3
"""Regenerate the cell ARGB fixture with openpyxl 3.1.5 (pip install openpyxl==3.1.5)."""
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.page import PageMargins

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'tests/ExcelRenderer.Tests/SampleInputs/CellArgb/10-cell-argb.xlsx'

wb = Workbook()
wb.properties.created = wb.properties.modified = datetime(2026, 10, 17)
ws = wb.active
ws.title = 'ARGB'
ws.sheet_view.showGridLines = False
ws.page_margins = PageMargins(left=0.25, right=0.25, top=0.25, bottom=0.25, header=0, footer=0)
ws.print_options.gridLines = False
ws.print_area = 'A1:E6'
for column, width in [('A', 18), ('B', 19), ('C', 19), ('D', 19), ('E', 19)]:
    ws.column_dimensions[column].width = width
for row in range(1, 7):
    ws.row_dimensions[row].height = 30
for col, label in enumerate(['ARGB prefix', 'Font', 'Fill', 'Border', 'Rich text'], 1):
    cell = ws.cell(1, col, label)
    cell.font = Font(name='Noto Sans JP', color='FF000000', size=11)
for row, prefix in enumerate(['00', '01', '80', 'FF', 'six-digit RGB'], 2):
    alpha = '' if prefix == 'six-digit RGB' else prefix
    ws.cell(row, 1, prefix).font = Font(name='Noto Sans JP', color='FF000000', size=11)
    ws.cell(row, 2, '日本語 ABC').font = Font(name='Noto Sans JP', color=alpha + '000000', size=11)
    ws.cell(row, 3).fill = PatternFill('solid', fgColor=alpha + '336699')
    side = Side(style='thick', color=alpha + 'CC3300')
    ws.cell(row, 4).border = Border(left=side, top=side, right=side, bottom=side)
    ws.cell(row, 5, CellRichText(TextBlock(InlineFont(rFont='Noto Sans JP', sz=11, color=alpha + '000000'), '日本語 ABC')))
    for col in range(2, 6):
        ws.cell(row, col).alignment = Alignment(horizontal='left', vertical='center')
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
buffer = io.BytesIO()
wb.save(buffer)
# Stable ZIP metadata makes regeneration reviewable. The six-digit row is
# deliberately passed through openpyxl, which prepends 00 to its RGB strings.
with ZipFile(buffer) as source, ZipFile(OUTPUT, 'w', ZIP_DEFLATED) as target:
    for name in source.namelist():
        info = ZipInfo(name, (2026, 10, 17, 0, 0, 0))
        info.compress_type = ZIP_DEFLATED
        data = source.read(name)
        if name == 'docProps/core.xml':
            # openpyxl updates modified on save even when a fixed property was supplied.
            data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]+',
                          rb'\g<1>2026-10-17T00:00:00Z', data)
        target.writestr(info, data)
print(OUTPUT)
